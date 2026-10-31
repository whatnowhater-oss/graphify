"""Done-for-you spreadsheet & contact list cleanup ($249 on Fiverr, up to 10,000 rows).

intake -> map_columns -> clean -> qa -> package -> deliver

Same pattern as X-ray with a different work step. When the customer's column
names differ from the cleaner's, the order waits for Claude to write a column
map instead of guessing. Buyer files are personal data: every copy is deleted
30 days after delivery, as the gig promises.
"""
from __future__ import annotations

import contextlib
import csv
import io
import json
import shutil
from pathlib import Path

from ..core import Order, Pipeline, Reject, Stage, Wait, ops_home
from . import clean_customers
from .common import deliver_stage, package_stage

EXPECTED = ["Customer Name", "Email", "Phone", "City", "State", "First Order", "Lifetime Spend", "Notes"]
MAX_ROWS = 10_000  # the gig's limit; bigger files get a custom offer


def _src(fields: dict) -> str:
    return fields.get("file_path") or fields.get("csv_path") or ""


def _validate(fields: dict) -> None:
    if not _src(fields):
        raise Reject("order has no file_path (the buyer's uploaded file)")
    if Path(_src(fields)).suffix.lower() not in (".csv", ".xlsx"):
        raise Reject("only .csv and .xlsx files are supported")


def _xlsx_to_csv(src: Path, dest: Path) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    with open(dest, "w", newline="") as fh:
        w = csv.writer(fh)
        for row in ws.iter_rows(values_only=True):
            w.writerow(["" if v is None else v for v in row])


def _csv_to_xlsx(src: Path, dest: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    with open(src, newline="") as fh:
        for row in csv.reader(fh):
            ws.append(row)
    wb.save(dest)


def intake(o: Order) -> None:
    src = Path(_src(o.fields))
    if not src.exists():
        raise Wait(f"buyer's file {src} not found; download it from the order and save it there", "jay")
    if src.stat().st_size > 50_000_000:
        raise Wait("file is over 50 MB; send the buyer a custom offer", "jay")
    if src.suffix.lower() == ".xlsx":
        _xlsx_to_csv(src, o.workdir / "original.csv")
        o.data["input_format"] = "xlsx"
    else:
        shutil.copy(src, o.workdir / "original.csv")
        o.data["input_format"] = "csv"
    with open(o.workdir / "original.csv", newline="", encoding="utf-8-sig") as fh:
        rows = sum(1 for _ in fh) - 1
    if rows > MAX_ROWS:
        raise Wait(f"{rows:,} rows, over the {MAX_ROWS:,}-row gig limit; send the buyer a custom offer", "jay")


def map_columns(o: Order) -> None:
    with open(o.workdir / "original.csv", newline="", encoding="utf-8-sig") as fh:
        header = next(csv.reader(fh))
    mapping_p = o.workdir / "column_map.json"
    if set(EXPECTED) <= set(header):
        mapping = {c: c for c in EXPECTED}
    elif mapping_p.exists():
        mapping = json.loads(mapping_p.read_text())
    else:
        raise Wait(f"columns {header} don't match {EXPECTED}; write {mapping_p} as "
                   f'{{"Expected Name": "Their Column"}} (leave out ones they lack), then `ops resume {o.id}`',
                   "claude")
    with open(o.workdir / "original.csv", newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    with open(o.workdir / "mapped.csv", "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=EXPECTED)
        w.writeheader()
        for r in rows:
            w.writerow({e: r.get(mapping.get(e, ""), "") for e in EXPECTED})
    o.data["rows_in"] = len(rows)


def clean(o: Order) -> None:
    clean_customers.changes.clear()
    with contextlib.redirect_stdout(io.StringIO()):
        clean_customers.main(str(o.workdir / "mapped.csv"), str(o.workdir / "clean.csv"),
                             str(o.workdir / "CHANGES_REPORT.md"))
    if o.data.get("input_format") == "xlsx":
        for name in ("clean", "clean_needs_review"):
            _csv_to_xlsx(o.workdir / f"{name}.csv", o.workdir / f"{name}.xlsx")


def qa(o: Order) -> None:
    with open(o.workdir / "clean.csv", newline="") as fh:
        out = list(csv.DictReader(fh))
    report = (o.workdir / "CHANGES_REPORT.md").read_text()
    problems = []
    if not out:
        problems.append("clean file is empty")
    if f"**Rows in:** {o.data['rows_in']}" not in report:
        problems.append("report row count doesn't match the input")
    merged = int(report.split("**Duplicates merged:** ")[1].split()[0])
    if len(out) + merged != o.data["rows_in"]:
        problems.append(f"rows lost: {o.data['rows_in']} in, {len(out)} out + {merged} merged")
    if problems:
        raise Wait("QA failed: " + "; ".join(problems), "claude")


def _files(o: Order) -> list[str]:
    ext = "xlsx" if o.data.get("input_format") == "xlsx" else "csv"
    return [f"clean.{ext}", f"clean_needs_review.{ext}", "CHANGES_REPORT.md"]


def _package(o: Order) -> None:
    package_stage(_files(o), lambda o: f"list-cleanup-{o.id[-6:]}").run(o)


def _customer_data(o: Order) -> list[Path]:
    paths = [o.home / "work" / o.id, o.home / "deliverables" / o.id]
    src = Path(_src(o.fields))
    if src.resolve().is_relative_to(ops_home()):
        paths.append(src)  # only files we stored; never touch a path outside the order desk
    return paths


def _body(o: Order) -> str:
    return (f"Hi {o.fields.get('name') or 'there'},\n\n"
            "Your cleaned list is attached. CHANGES_REPORT.md explains every fix. The needs-review file "
            "lists the rows where a guess would have been worse than a blank, so you can decide those.\n\n"
            "Your files are deleted from our side 30 days after delivery.\n\n"
            "Thanks for your order,\n{sender}\n")


PIPELINE = Pipeline(
    product="cleanup", title="List cleanup", price_usd=249, validate=_validate,
    turnaround_hours=48, retention_days=30, customer_data=_customer_data,
    stages=[
        Stage("intake", intake, "find the buyer's file, convert Excel, check the row limit"),
        Stage("map_columns", map_columns, "match their columns to the cleaner's"),
        Stage("clean", clean, "normalize, merge duplicates, flag problems"),
        Stage("qa", qa, "no rows lost, report matches the data"),
        Stage("package", _package, "zip the deliverable"),
        deliver_stage(lambda o: "Your cleaned customer list", _body),
    ],
)
